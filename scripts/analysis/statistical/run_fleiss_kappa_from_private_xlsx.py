"""Calculate reviewer agreement from a private scoring workbook.

The workbook contains individual reviewer ratings and must remain outside the
public repository. This command validates the original 48-by-11 scoring form,
checks every form row against a stable title/DOI contract, and writes aggregate
agreement results for all 48 records scored by the 11 reviewers. Three records
later excluded from the scientific synthesis remain in this reliability
analysis because they were part of the completed scoring exercise.

Example:
    python scripts/analysis/statistical/run_fleiss_kappa_from_private_xlsx.py \
        --input-xlsx C:/private/RECEIVED_SCORES.xlsx \
        --output-dir tables
"""

from __future__ import annotations

import argparse
import csv
import json
import unicodedata
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
from openpyxl import load_workbook
from statsmodels.stats.inter_rater import fleiss_kappa


EXPECTED_ITEMS = 48
EXPECTED_RATERS = 11
REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_CANONICAL_DATA = REPO_ROOT / "data" / "data-clean.csv"
DEFAULT_SCORING_ORDER = REPO_ROOT / "data" / "reviewer_scoring_order.csv"
SUMMARY_NAME = "analysis_fleiss_kappa_summary.csv"
ITEM_NAME = "analysis_fleiss_kappa_item_agreement.csv"


def _as_int(value: object, *, paper: int, reviewer: int, scale: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(
            f"Missing or non-numeric {scale} score for paper {paper}, "
            f"reviewer {reviewer}"
        )
    if int(value) != value:
        raise ValueError(f"Non-integer {scale} score for paper {paper}, reviewer {reviewer}")
    return int(value)


def _normalize_title(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKC", text)
    text = "".join(char for char in text if unicodedata.category(char) != "Cf")
    text = text.translate(str.maketrans({"–": "-", "—": "-", "’": "'", "‘": "'"}))
    return " ".join(text.replace("\u00a0", " ").split()).casefold()


def _canonical_studies(path: Path):
    import pandas as pd

    if not path.exists():
        raise FileNotFoundError(f"Canonical study data not found: {path}")
    frame = pd.read_csv(path)
    required = {"Paper_ID", "Title", "DOI"}
    if not required.issubset(frame.columns):
        raise ValueError(f"Canonical study data must contain {sorted(required)}")
    frame = frame.copy()
    frame["Paper_ID"] = pd.to_numeric(frame["Paper_ID"], errors="raise").astype(int)
    if frame.empty or frame["Paper_ID"].tolist() != list(range(1, len(frame) + 1)):
        raise ValueError("Canonical study data must have contiguous Paper_ID values from 1")
    titles = [_normalize_title(value) for value in frame["Title"]]
    dois = [_normalize_doi(value) for value in frame["DOI"]]
    if not all(titles) or len(set(titles)) != len(frame):
        raise ValueError("Canonical study titles must be non-empty and unique")
    if not all(dois) or len(set(dois)) != len(frame):
        raise ValueError("Canonical study DOIs must be non-empty and unique")
    frame["_normalized_title"] = titles
    frame["_normalized_doi"] = dois
    return frame


def _normalize_doi(value: object) -> str:
    text = "" if value is None else str(value).strip().casefold()
    text = text.removeprefix("https://doi.org/").removeprefix("http://doi.org/")
    return text.split("?", 1)[0].rstrip("/ ")


def _scoring_contract(path: Path, canonical_data: Path):
    import pandas as pd

    if not path.exists():
        raise FileNotFoundError(f"Reviewer scoring-order contract not found: {path}")
    contract = pd.read_csv(path)
    required = {
        "Form_Paper_ID", "Title", "DOI", "Include_In_Final_Corpus", "Final_Paper_ID"
    }
    if not required.issubset(contract.columns):
        raise ValueError(f"Scoring-order contract must contain {sorted(required)}")
    if len(contract) != EXPECTED_ITEMS:
        raise ValueError(f"Scoring-order contract must contain {EXPECTED_ITEMS} form rows")
    form_ids = pd.to_numeric(contract["Form_Paper_ID"], errors="raise").astype(int)
    if form_ids.tolist() != list(range(1, EXPECTED_ITEMS + 1)):
        raise ValueError("Form_Paper_ID must be contiguous 1-48 in scoring order")
    contract["_normalized_title"] = contract["Title"].map(_normalize_title)
    contract["_normalized_doi"] = contract["DOI"].map(_normalize_doi)
    if contract["_normalized_title"].eq("").any() or contract["_normalized_title"].duplicated().any():
        raise ValueError("Scoring-order titles must be unique and non-empty")
    if contract["_normalized_doi"].eq("").any() or contract["_normalized_doi"].duplicated().any():
        raise ValueError("Scoring-order DOIs must be unique and non-empty")
    include = contract["Include_In_Final_Corpus"].map(_as_bool)
    final_ids = pd.to_numeric(contract.loc[include, "Final_Paper_ID"], errors="raise").astype(int)
    if final_ids.tolist() != list(range(1, int(include.sum()) + 1)):
        raise ValueError("Included Final_Paper_ID values must be contiguous in form order")

    canonical = _canonical_studies(canonical_data)
    active = contract.loc[include].sort_values("Final_Paper_ID")
    if active["_normalized_title"].tolist() != canonical["_normalized_title"].tolist():
        raise ValueError("Scoring-order eligible titles do not match canonical corpus order")
    if active["_normalized_doi"].tolist() != canonical["_normalized_doi"].tolist():
        raise ValueError("Scoring-order eligible DOIs do not match canonical corpus order")
    contract["_include"] = include
    return contract, canonical


def _as_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    if text in {"true", "1", "yes"}:
        return True
    if text in {"false", "0", "no"}:
        return False
    raise ValueError(f"Invalid Include_In_Final_Corpus value: {value!r}")


def _read_matrix(
    path: Path,
    canonical_data: Path | None = None,
    *,
    include_titles: bool = False,
    include_excluded_items: bool = False,
    scoring_order_path: Path | None = None,
) -> tuple[np.ndarray, np.ndarray] | tuple[np.ndarray, np.ndarray, list[str]]:
    if not path.exists():
        raise FileNotFoundError(f"Private reviewer workbook not found: {path}")

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) != EXPECTED_ITEMS + 2:
        raise ValueError(
            f"Expected {EXPECTED_ITEMS} papers plus two header rows; found {len(rows)} rows"
        )
    if len(rows[0]) < 3 or rows[0][0] is not None or rows[1][0] != "PAPER":
        raise ValueError("Unexpected reviewer workbook header")

    score_columns: list[tuple[int, int, str]] = []
    for lm_col in range(3, len(rows[1]), 2):
        tr_col = lm_col + 1
        if tr_col >= len(rows[1]):
            break
        reviewer = rows[0][lm_col]
        lm_header = rows[1][lm_col]
        tr_header = rows[1][tr_col]
        if not reviewer or "LMIC" not in str(lm_header) or "TR" not in str(tr_header):
            continue
        score_columns.append((lm_col, tr_col, str(reviewer)))

    if len(score_columns) != EXPECTED_RATERS:
        raise ValueError(
            f"Expected {EXPECTED_RATERS} reviewer pairs; found {len(score_columns)}"
        )

    canonical_path = canonical_data or DEFAULT_CANONICAL_DATA
    contract_path = scoring_order_path or DEFAULT_SCORING_ORDER
    contract, canonical = _scoring_contract(contract_path, canonical_path)

    lmic: list[list[int]] = []
    tr: list[list[int]] = []
    selected_titles: list[str] = []
    for row_offset, row in enumerate(rows[2:], start=1):
        paper = row[0]
        if paper != row_offset:
            raise ValueError(f"Paper rows must be ordered 1-{EXPECTED_ITEMS}; found {paper}")
        contract_row = contract.iloc[row_offset - 1]
        workbook_title = _normalize_title(row[1])
        if workbook_title != contract_row["_normalized_title"]:
            raise ValueError(
                f"Paper {row_offset} title does not match reviewer form order: "
                f"{row[1]!r}"
            )
        row_lmic = [
            _as_int(row[lm_col], paper=row_offset, reviewer=i, scale="LMIC")
            for i, (lm_col, _, _) in enumerate(score_columns, start=1)
        ]
        row_tr = [
            _as_int(row[tr_col], paper=row_offset, reviewer=i, scale="TR")
            for i, (_, tr_col, _) in enumerate(score_columns, start=1)
        ]
        if not all(1 <= score <= 5 for score in row_lmic):
            raise ValueError("LMIC scores must be integers from 1 to 5")
        if not all(0 <= score <= 5 for score in row_tr):
            raise ValueError("TR scores must be integers from 0 to 5")

        if _as_bool(contract_row["_include"]) or include_excluded_items:
            lmic.append(row_lmic)
            tr.append(row_tr)
            selected_titles.append(str(contract_row["Title"]))

    lmic_array = np.asarray(lmic, dtype=int)
    tr_array = np.asarray(tr, dtype=int)
    if not np.all((1 <= lmic_array) & (lmic_array <= 5)):
        raise ValueError("LMIC scores must be integers from 1 to 5")
    if not np.all((0 <= tr_array) & (tr_array <= 5)):
        raise ValueError("TR scores must be integers from 0 to 5")
    expected_items = EXPECTED_ITEMS if include_excluded_items else len(canonical)
    if lmic_array.shape != (expected_items, EXPECTED_RATERS):
        raise ValueError("Reviewer rating dimensions do not match the selected study set")
    if include_titles:
        return lmic_array, tr_array, selected_titles
    return lmic_array, tr_array


def _agreement_rows(
    ratings: np.ndarray,
    categories: Sequence[int],
    analysis: str,
    titles: Sequence[str] | None = None,
) -> tuple[dict[str, object], list[dict[str, object]]]:
    n_items, n_raters = ratings.shape
    counts = np.asarray(
        [[int(np.sum(row == category)) for category in categories] for row in ratings],
        dtype=int,
    )
    kappa = float(fleiss_kappa(counts, method="fleiss"))
    pair_agreement = (np.sum(counts**2, axis=1) - n_raters) / (
        n_raters * (n_raters - 1)
    )
    majority = counts.max(axis=1) / n_raters
    p_bar = float(pair_agreement.mean())
    p_e = float(np.square(counts.sum(axis=0) / (n_items * n_raters)).sum())
    summary = {
        "Analysis": analysis,
        "Items": n_items,
        "Raters": n_raters,
        "Fleiss_kappa": kappa,
        "Observed_agreement_P_bar": p_bar,
        "Expected_agreement_P_e": p_e,
        "Mean_majority_share": float(majority.mean()),
        "Unanimous_items": int(np.sum(counts.max(axis=1) == n_raters)),
        "Category_counts": "; ".join(
            f"{category}={int(counts[:, index].sum())}"
            for index, category in enumerate(categories)
        ),
    }
    item_rows = []
    for index, row_counts in enumerate(counts, start=1):
        max_count = int(row_counts.max())
        modal = [
            str(category)
            for category, count in zip(categories, row_counts)
            if count == max_count
        ]
        item = {
                "Analysis": analysis,
                "Scoring_Form_ID": index,
                "Mean_pairwise_agreement": float(pair_agreement[index - 1]),
                "Majority_share": float(majority[index - 1]),
                "Modal_score": "/".join(modal),
                "Unanimous": bool(max_count == n_raters),
        }
        if titles is not None:
            item["Title"] = titles[index - 1]
        item_rows.append(item)
    return summary, item_rows


def _write_csv(path: Path, rows: Iterable[dict[str, object]]) -> None:
    rows = list(rows)
    if not rows:
        raise ValueError(f"Cannot write empty output: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-xlsx", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--canonical-data", type=Path, default=DEFAULT_CANONICAL_DATA)
    parser.add_argument("--scoring-order", type=Path, default=DEFAULT_SCORING_ORDER)
    args = parser.parse_args()

    lmic, tr, titles = _read_matrix(
        args.input_xlsx.resolve(),
        args.canonical_data.resolve(),
        include_titles=True,
        include_excluded_items=True,
        scoring_order_path=args.scoring_order.resolve(),
    )
    summaries = []
    item_rows = []
    for name, ratings, categories in (
        ("LMIC_Relevance_Score", lmic, list(range(1, 6))),
        ("TR_Score", tr, list(range(0, 6))),
    ):
        summary, items = _agreement_rows(ratings, categories, name, titles)
        summaries.append(summary)
        item_rows.extend(items)

    output_dir = args.output_dir.resolve()
    summary_path = output_dir / SUMMARY_NAME
    item_path = output_dir / ITEM_NAME
    _write_csv(summary_path, summaries)
    _write_csv(item_path, item_rows)
    print(
        json.dumps(
            {
                "items": int(lmic.shape[0]),
                "form_items": EXPECTED_ITEMS,
                "final_eligible_items": int(len(_canonical_studies(args.canonical_data.resolve()))),
                "excluded_from_final_synthesis_after_scoring": EXPECTED_ITEMS - int(len(_canonical_studies(args.canonical_data.resolve()))),
                "raters": EXPECTED_RATERS,
                "summary": str(summary_path),
                "item_agreement": str(item_path),
                "results": summaries,
                "individual_ratings_written": False,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
